"""Planilhas sintéticas que reproduzem as manhas das planilhas reais.

As planilhas de produção não são versionadas (trazem nomes de funcionários e
programação não veiculada), então os testes que dependiam delas pulavam num
clone limpo — justamente os que provam os defeitos críticos.

Estas fixtures recriam, de propósito, cada característica do formato real que
quebrou o sistema:

ESCALA INDIVIDUAL
  * `'Nome '` com espaço no fim e `'Início'`/`'Pré'` acentuados
  * `Data`, `Início`, `Fim` como datetime (não string) — era o que virava epoch
    no `to_json()` e destruía a chave de comparação
  * coluna `Data_raw`, que fazia o achatamento de blocos confundir a escala com
    uma grade de TV e descartar a coluna `Nome`
  * evento em `Atividade/Descrição`, não em `Evento` — o fallback do código
    apontava para `'Evento/Programa'`, que nunca existe
  * folgas com o texto em `Descrição` e `Atividade/Descrição` vazio
  * `Local de locução` e `Local` lado a lado, com valores diferentes
  * `Plataforma` como `Sportv 2`

GRADE DE TV DENSA
  * cabeçalho fora da primeira linha
  * colunas de contexto à esquerda (`DATA GRADE`) e 4 blocos lado a lado
  * `DATA` de cada bloco vazia: a data da linha mora só no prefixo
"""

from __future__ import annotations

import datetime as dt
import os
from typing import Dict, List, Optional

import pandas as pd

# ---------------------------------------------------------------------------
# Dados base — 5 pessoas, 21 a 25/09/2026
# ---------------------------------------------------------------------------

ANO, MES = 2026, 9

#: (funcionario, dia, hora_ini, hora_fim, evento, plataforma, elenco)
ESCALA_PADRAO = [
    ("Marina Tavares",  21, (18, 30), (20, 0),  "SAO PAULO X PALMEIRAS", "Sportv 2", "Marina Tavares ; Rui Campos ; Sofia Dantas"),
    ("Marina Tavares",  23, (11, 30), (12, 0),  "COPA DO MUNDO FEMININA SUB-20", "Sportv 3", "Marina Tavares ; Rui Campos"),
    ("Rui Campos",      21, (23, 30), (1, 0),   "TROCA DE PASSES", "Sportv", "Rui Campos ; Carlos Vieira"),
    ("Rui Campos",      25, (15, 45), (17, 45), "HUNGRIA X UCRANIA", "Sportv 2", "Rui Campos"),
    ("Carlos Vieira",   23, (11, 30), (12, 0),  "COPA DO MUNDO FEMININA SUB-20", "Sportv 2", "Carlos Vieira ; Rui Campos ; Sofia Dantas"),
    ("Sofia Dantas",    22, (17, 0),  (18, 30), "GRAND SLAM DE JUDO", "Sportv", "Sofia Dantas"),
    ("Paulo Reis",      24, (20, 0),  (22, 0),  "FLUMINENSE X PRAIA CLUBE", "Premiere 3", "Paulo Reis"),
]

#: (funcionario, dia, descricao) — folga: evento vazio, 00:00 às 00:00
FOLGAS_PADRAO = [
    ("Marina Tavares", 22, "Day Off / Folga"),
    ("Rui Campos",     24, "Comp Day / Folga Compensatória"),
    ("Carlos Vieira",  25, "Comp Day / Folga Compensatória"),
]

#: Eventos da grade de TV. (dia, hora, evento) — casa com a escala acima.
GRADE_PADRAO = [
    (21, (18, 30), "SAO PAULO X PALMEIRAS"),
    (21, (23, 30), "TROCA DE PASSES"),
    (21, (7, 0),   "TROCA DE PASSES"),          # 2ª exibição: testa desempate
    (22, (17, 0),  "GRAND SLAM DE JUDO"),
    (23, (11, 30), "COPA DO MUNDO FEMININA SUB-20 DA FIFA 2026 - SEMIFINAL"),
    (24, (20, 0),  "FLUMINENSE X PRAIA CLUBE"),
    (25, (15, 45), "HUNGRIA X UCRANIA"),
]


def _ts(dia: int, hora=(0, 0)) -> dt.datetime:
    return dt.datetime(ANO, MES, dia, hora[0] % 24, hora[1])


# ---------------------------------------------------------------------------
# Escala individual
# ---------------------------------------------------------------------------

def criar_escala(
    caminho: str,
    escala: Optional[List] = None,
    folgas: Optional[List] = None,
) -> str:
    """Grava uma escala no formato real e devolve o caminho."""
    escala = ESCALA_PADRAO if escala is None else escala
    folgas = FOLGAS_PADRAO if folgas is None else folgas

    registros: List[Dict] = []
    for nome, dia, ini, fim, evento, plataforma, elenco in escala:
        registros.append({
            "Nome ": nome,
            "Tipo Atividade ": "Booking",
            "Descrição": "",
            "Data": _ts(dia, ini),
            "Dia": _ts(dia, ini),
            "Pré": float("nan"),
            "Início": _ts(dia, ini),
            "Fim": _ts(dia, fim),
            "Event Group": evento.split(" - ")[0],
            "Atividade/Descrição": evento,
            "Equipe": "Esp - Talentos",
            "Plataforma": plataforma,
            "Local de locução": "Offtube ION",
            "Local": "Internacional - Polônia",
            "Narrador": "",
            "Comentarista": "",
            "Elenco": elenco,
            "Coordenador": "COORDENADOR TESTE",
            "Produtor": "PRODUTOR TESTE",
            "Data_raw": _ts(dia, ini),
        })

    for nome, dia, descricao in folgas:
        registros.append({
            "Nome ": nome,
            "Tipo Atividade ": "Other Time Off",
            "Descrição": descricao,
            "Data": _ts(dia),
            "Dia": _ts(dia),
            "Pré": float("nan"),
            "Início": _ts(dia),
            "Fim": _ts(dia),
            "Event Group": "",
            "Atividade/Descrição": "",
            "Equipe": "Esp - Talentos",
            "Plataforma": "",
            "Local de locução": "",
            "Local": "",
            "Narrador": "",
            "Comentarista": "",
            "Elenco": "",
            "Coordenador": "",
            "Produtor": "",
            "Data_raw": _ts(dia),
        })

    os.makedirs(os.path.dirname(os.path.abspath(caminho)), exist_ok=True)
    pd.DataFrame(registros).to_excel(caminho, index=False)
    return caminho


# ---------------------------------------------------------------------------
# Grade de TV densa
# ---------------------------------------------------------------------------

_PREFIXO = ["DATA GRADE", "MÊS", "DIA DA SEMANA", "DATA DA SEMANA REAL",
            "DATA REAL", "FAIXA HORÁRIA", "SEMANA"]
_BLOCO = ["DATA", "HORA", "", "EVENTO/PROGRAMA", "OBSERVAÇÃO"]
_DIAS_ABREV = ["SEG", "TER", "QUA", "QUI", "SEX", "SAB", "DOM"]


def criar_grade_tv(
    caminho: str,
    eventos: Optional[List] = None,
    blocos: int = 4,
    data_no_bloco: bool = False,
    linha_cabecalho: int = 3,
) -> str:
    """Grava uma grade de TV densa e devolve o caminho.

    Com `data_no_bloco=False` (o padrão, igual ao arquivo real) a coluna `DATA`
    de cada bloco fica vazia e a data só existe no prefixo da linha. Era isso
    que deixava 98% das linhas sem data após o achatamento.
    """
    import openpyxl

    eventos = GRADE_PADRAO if eventos is None else eventos

    wb = openpyxl.Workbook()
    aba = wb.active
    aba.title = f"SETEMBRO {ANO} - CANAIS SPORTV"

    aba.cell(row=1, column=1, value=f"GRADE DE EVENTOS {ANO}")

    cabecalho = list(_PREFIXO) + list(_BLOCO) * blocos
    for col, titulo in enumerate(cabecalho, start=1):
        if titulo:
            aba.cell(row=linha_cabecalho, column=col, value=titulo)

    # Um evento por linha, distribuído entre os blocos (como os canais da grade).
    linha = linha_cabecalho + 1
    for i, (dia, hora, evento) in enumerate(eventos):
        data = _ts(dia)
        aba.cell(row=linha, column=1, value=data)
        aba.cell(row=linha, column=2, value=MES)
        aba.cell(row=linha, column=3, value=_DIAS_ABREV[data.weekday()])
        aba.cell(row=linha, column=4, value=_DIAS_ABREV[data.weekday()])
        aba.cell(row=linha, column=5, value=data)
        aba.cell(row=linha, column=6, value="MAT" if hora[0] < 12 else "NOI")
        aba.cell(row=linha, column=7, value=35)

        base = len(_PREFIXO) + (i % blocos) * len(_BLOCO) + 1
        if data_no_bloco:
            aba.cell(row=linha, column=base, value=data)
        aba.cell(row=linha, column=base + 1, value=dt.time(hora[0] % 24, hora[1]))
        aba.cell(row=linha, column=base + 3, value=evento)
        linha += 1

    os.makedirs(os.path.dirname(os.path.abspath(caminho)), exist_ok=True)
    wb.save(caminho)
    wb.close()
    return caminho


def criar_grade_de_outro_periodo(caminho: str) -> str:
    """Grade cujas datas não tocam a escala padrão (mês seguinte)."""
    global MES
    original = MES
    try:
        MES = 10
        return criar_grade_tv(caminho, eventos=[
            (5, (20, 0), "CAMPEONATO DE OUTUBRO"),
            (6, (21, 0), "OUTRO EVENTO DE OUTUBRO"),
        ])
    finally:
        MES = original


def criar_grade_especializada(caminho: str) -> str:
    """Grade de recorte estreito: cobre as datas, mas quase nenhum evento.

    É o caso da grade de Combate sobre uma escala de futebol e vôlei: cancelar
    o que ela não menciona seria errado.
    """
    return criar_grade_tv(caminho, eventos=[
        (21, (22, 0), "UFC FIGHT NIGHT"),
        (22, (23, 0), "JUNGLE FIGHT 144"),
        (23, (21, 0), "ONE FIGHT NIGHT 45"),
        (24, (22, 0), "UFC BJJ 11"),
        (25, (23, 0), "PESAGEM JUNGLE FIGHT"),
    ], blocos=1)
